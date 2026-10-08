"""Builds the formula-driven Gantt workbooks.

    python3 build_gantt_template.py template  Gantt_Chart_Template.xlsx
    python3 build_gantt_template.py execution Project_Execution_Gantt.xlsx

Features: weekly-off setting, holiday list, FS / SS / FF links with lag, Day / Week / Month view,
status tracking with overshoot + knock-on delay propagation.
"""
import datetime as dt
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.formatting.rule import Rule, DataBarRule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.comments import Comment

KIND = sys.argv[1] if len(sys.argv) > 1 else "template"
OUT = sys.argv[2] if len(sys.argv) > 2 else ("Project_Execution_Gantt.xlsx" if KIND == "execution" else "Gantt_Chart_Template.xlsx")
EXEC = KIND == "execution"

FIRST, LAST = 11, 310            # task rows (300 tasks ready; copy last row down for more)
NCOL, TL0 = 120, 15              # timeline columns, first timeline column (O)
TLN = TL0 + NCOL - 1
INFO = "'Project Info'!"
PSD, TGT, SD = INFO + "$C$9", INFO + "$C$10", INFO + "$C$11"
WKN = INFO + "$H$12"             # weekend code derived from Weekly Off setting
HOL = "Holidays!$A$6:$A$205"
BANNER = INFO + "$B$15"

# ---- palette -------------------------------------------------------------
NAVY, BLUE, TEAL, PURPLE, GREY = "1F3864", "2F5597", "1B7F79", "6A3D9A", "595959"
C_DONE, C_PROG, C_NOT, C_HOLD, C_OVER, C_AFF, C_GHOST = (
    "2E9E5B", "2F80ED", "A9C1E8", "F2C94C", "E53935", "F2994A", "D9D9D9")
C_WEEKEND, C_HOLIDAY, C_STATUS = "E4E4E4", "F8C8DC", "FFE699"
INPUT = "FFFBE6"
FN = "Arial"


def font(sz=10, b=False, c="000000", i=False):
    return Font(name=FN, size=sz, bold=b, color=c, italic=i)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


thin = Side(style="thin", color="D0D0D0")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def put(ws, ref, v=None, f=None, fl=None, al=None, nf=None, bd=None):
    c = ws[ref]
    if v is not None:
        c.value = v
    c.font = f or font()
    if fl:
        c.fill = fill(fl)
    if al:
        c.alignment = al
    if nf:
        c.number_format = nf
    if bd:
        c.border = bd
    return c


def dxf(bg=None, fc=None, bold=False, nf=None, nfid=None):
    kw = {}
    if bg:
        kw["fill"] = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
    if fc or bold:
        kw["font"] = Font(color=fc, bold=bold)
    if nf:
        kw["numFmt"] = NumberFormat(numFmtId=nfid, formatCode=nf)
    return DifferentialStyle(**kw)


def frule(ws, rng, formula, **kw):
    stop = kw.pop("stop", False)
    ws.conditional_formatting.add(rng, Rule(type="expression", formula=[formula], dxf=dxf(**kw), stopIfTrue=stop))


def addv(ws, dv, rng):
    ws.add_data_validation(dv)
    dv.add(rng)


def W(a, n):                      # WORKDAY honouring weekly-off + holidays
    return f"WORKDAY.INTL({a},{n},{WKN},{HOL})"


def ND(a, b):                     # NETWORKDAYS honouring weekly-off + holidays
    return f"NETWORKDAYS.INTL({a},{b},{WKN},{HOL})"


# =========================================================================
# DATA
# =========================================================================
if EXEC:
    PROJECT = ("1.8M C-Band Antenna System – 500 Sets", "<Client name>", "<Project manager>", "PRJ-2026-001")
    START, TARGET = dt.date(2026, 10, 15), dt.date(2027, 4, 23)
    STATUS_DATE = "=TODAY()"
    WEEKLY_OFF = "Sunday only"
    HOLIDAYS = [(dt.date(2026, 10, 20), "Dussehra"),
                (dt.date(2026, 11, 7), "Diwali – day 1 (confirm date)"),
                (dt.date(2026, 11, 9), "Diwali – day 2 (confirm date)"),
                (dt.date(2026, 11, 10), "Diwali – day 3 (confirm date)"),
                (dt.date(2026, 11, 11), "Diwali – day 4 (confirm date)"),
                (dt.date(2027, 1, 1), "New Year")]
    COMPONENTS = ["REFLECTOR", "FEED", "METAL"]
    FS_ = (None, "FS", 0)
    # (name, owner, pred1, pred2, duration, basis)
    TASKS = [
        ("REFLECTOR | Finalize 1.8M mold design", "Design", None, None, 7, "Given: 7 days"),
        ("REFLECTOR | Mold manufacturing", "Tooling", (1, "FS", 0), None, 70, "Given: 70 days"),
        ("REFLECTOR | Reflector validation", "QA", (2, "FS", 0), None, 7, "Given: 7 days"),
        ("REFLECTOR | Production (500 nos @ 35/day)", "Production", (3, "FS", 0), None, 15, "500 ÷ 35/day = 14.3 → 15 days; starts after validation"),
        ("REFLECTOR | Finishing & packing (35/day)", "Finishing", (4, "FS", 0), None, 15, "500 ÷ 35/day = 14.3 → 15 days. No overlap with production; starts next working day = ≥12 h cooling gap"),
        ("REFLECTOR | Dispatch – 8 vehicles (rolling)", "Logistics", (5, "SS", 2), (5, "FF", 1), 14,
         "70 nos/vehicle → 7×70 + 1×10. First 70 packed after 2 days of finishing (start = finishing start +2); last vehicle leaves 1 day after packing ends (end = finishing end +1). 15−2+1 = 14 days"),
        ("FEED | Finalize mold design (6 components)", "Design", None, None, 15, "Given: 15 days (Feed Horn, Splitter, Tx/Rx Filter, Tx/Rx OMT)"),
        ("FEED | Mold making – Round 1 (Horn, Splitter, Tx Filter)", "Tooling", (7, "FS", 0), None, 30, "Given: 30 days, 3 molds in parallel"),
        ("FEED | Mold making – Round 2 (Rx Filter, Tx/Rx OMT)", "Tooling", (8, "FS", 0), None, 30, "Given: next 3 molds, 30 days"),
        ("FEED | Molding – Round-1 components (1,500 nos)", "Molding", (8, "FS", 0), None, 4, "3 types × 500 = 1,500 ÷ 400/day (2 machines × 200) = 3.75 → 4 days"),
        ("FEED | Molding – Round-2 components (1,500 nos)", "Molding", (9, "FS", 0), (10, "FS", 0), 4, "1,500 ÷ 400/day = 3.75 → 4 days; needs round-2 molds and free machines"),
        ("FEED | Machining – Round 1, Batch 1 (750 nos)", "Machining", (10, "FS", 0), None, 15, "250 sets × 3 types = 750 ÷ 50/day (2 machines × 25) = 15 days"),
        ("FEED | Machining – Round 1, Batch 2 (750 nos)", "Machining", (12, "FS", 0), None, 15, "750 ÷ 50/day = 15 days; done before round-2 molds are ready to avoid idle machines"),
        ("FEED | Machining – Round 2, Batch 1 (750 nos)", "Machining", (11, "FS", 0), (13, "FS", 0), 15, "750 ÷ 50/day = 15 days; completes the 250-set batch 1"),
        ("FEED | Machining – Round 2, Batch 2 (750 nos)", "Machining", (14, "FS", 0), None, 15, "750 ÷ 50/day = 15 days; completes batch 2"),
        ("FEED | Transport Batch 1 (250 sets) to finishing site", "Logistics", (14, "FS", 0), None, 1, "Assumed 1 day"),
        ("FEED | Finishing – Batch 1 (1,500 components)", "Finishing", (16, "FS", 0), None, 25, "1,500 ÷ 60/day (2 stations × 30) = 25 days"),
        ("FEED | Assembly & testing – Batch 1 (250 sets)", "Assembly & QA", (17, "FS", 0), None, 9, "250 ÷ 30/day = 8.3 → 9 days"),
        ("FEED | Packing – Batch 1 (100/day)", "Packing", (18, "FS", 0), None, 3, "250 ÷ 100/day = 2.5 → 3 days"),
        ("FEED | Dispatch – Vehicle 1 (150 sets)", "Logistics", (19, "SS", 2), None, 1, "150 sets packed after 1.5 days → leaves after 2 packing days (start = packing start +2)"),
        ("FEED | Transport Batch 2 (250 sets)", "Logistics", (15, "FS", 0), None, 1, "Assumed 1 day"),
        ("FEED | Finishing – Batch 2 (1,500 components)", "Finishing", (21, "FS", 0), (17, "FS", 0), 25, "1,500 ÷ 60/day = 25 days; finishing stations busy with batch 1 first"),
        ("FEED | Assembly & testing – Batch 2 (250 sets)", "Assembly & QA", (22, "FS", 0), (18, "FS", 0), 9, "250 ÷ 30/day = 8.3 → 9 days"),
        ("FEED | Packing – Batch 2 (100/day)", "Packing", (23, "FS", 0), (19, "FS", 0), 3, "250 ÷ 100/day = 2.5 → 3 days"),
        ("FEED | Dispatch – Vehicles 2-4 (rolling)", "Logistics", (24, "SS", 1), (24, "FF", 1), 3,
         "V2 = 100 left from batch 1 + 50 from batch 2 (ready after 1 packing day), V3 = 150, V4 = last 50; last leaves 1 day after packing ends. Day 2 to day 4 = 3 days"),
        ("METAL | Finalize manufacturing drawings", "Design", None, None, 12, "Given: 12 days after PO"),
        ("METAL | Manufacture all components", "Fabrication", (26, "FS", 0), None, 45, "Given: 45 days"),
        ("METAL | Powder coating all components", "Coating", (27, "FS", 0), None, 10, "Given: 10 days"),
        ("METAL | Assembly & packing (75 sets/day)", "Assembly", (28, "FS", 0), None, 7, "500 ÷ 75/day = 6.7 → 7 days"),
        ("METAL | Dispatch – 10 vehicles (rolling)", "Logistics", (29, "SS", 1), (29, "FF", 1), 7,
         "50 sets/vehicle → 10 vehicles; first 50 packed after 1 day (start = packing start +1); last leaves 1 day after packing ends. Day 2 to day 8 = 7 days"),
    ]
    EX_TASKS = [dict(name=n, owner=o, p1=p1, p2=p2, dur=d) for n, o, p1, p2, d, _ in TASKS]
else:
    PROJECT = ("Website Redesign & Launch", "Acme Corporation", "Your Name", "PRJ-2026-001")
    START, TARGET = dt.date(2026, 10, 5), dt.date(2026, 12, 18)
    STATUS_DATE = dt.date(2026, 11, 6)
    WEEKLY_OFF = "Saturday & Sunday"
    HOLIDAYS = [(dt.date(2026, 10, 20), "Sample holiday 1 (edit me)"), (dt.date(2026, 11, 10), "Sample holiday 2 (edit me)"),
                (dt.date(2026, 12, 25), "Christmas Day")]
    COMPONENTS = ["DESIGN", "BUILD", "LAUNCH"]
    D = dt.date
    # name, owner, p1, p2, dur, snet, status, act start, act end, pct, revised
    EX = [
        ("DESIGN | Project kickoff & charter", "Project Manager", None, None, 2, None, "Completed", D(2026, 10, 5), D(2026, 10, 6), 1, None),
        ("DESIGN | Requirements gathering", "Business Analyst", 1, None, 5, None, "Completed", D(2026, 10, 7), D(2026, 10, 15), 1, None),
        ("DESIGN | Technical architecture", "Architect", 2, None, 4, None, "Completed", D(2026, 10, 16), D(2026, 10, 21), 1, None),
        ("DESIGN | UX wireframes", "UX Designer", 2, None, 10, None, "In Progress", D(2026, 10, 16), None, 0.7, 17),
        ("DESIGN | Visual design", "UI Designer", 4, None, 8, None, "Not Started", None, None, None, None),
        ("BUILD | Database setup", "DBA", 3, None, 5, None, "Completed", D(2026, 10, 22), D(2026, 10, 30), 1, None),
        ("BUILD | Backend API development", "Dev Team A", 6, None, 12, None, "In Progress", D(2026, 11, 2), None, 0.3, None),
        ("BUILD | Frontend development", "Dev Team B", 5, 6, 10, None, "Not Started", None, None, None, None),
        ("BUILD | Content creation", "Content Lead", 2, None, 10, D(2026, 10, 26), "In Progress", D(2026, 10, 26), None, 0.8, None),
        ("LAUNCH | Integration & QA testing", "QA Team", 7, 8, 8, None, "Not Started", None, None, None, None),
        ("LAUNCH | User acceptance testing", "Client", 10, None, 5, None, "Not Started", None, None, None, None),
        ("LAUNCH | Training & documentation", "Trainer", 9, None, 5, None, "On Hold", None, None, None, None),
        ("LAUNCH | Deployment / go-live", "DevOps", 11, 12, 2, None, "Not Started", None, None, None, None),
        ("LAUNCH | Project closure & handover", "Project Manager", 13, None, 2, None, "Not Started", None, None, None, None),
    ]
    EX_TASKS = []
    for n, o, p1, p2, d, sn, st, a1, a2, pc, rv in EX:
        EX_TASKS.append(dict(name=n, owner=o, p1=(p1, "FS", 0) if p1 else None, p2=(p2, "FS", 0) if p2 else None,
                             dur=d, snet=sn, status=st, acts=a1, acte=a2, pct=pc, rev=rv))

wb = Workbook()
wi = wb.active
wi.title = "Project Info"
wt = wb.create_sheet("Tasks")
wg = wb.create_sheet("Gantt")
wh = wb.create_sheet("Holidays")
wa = wb.create_sheet("Assumptions") if EXEC else None
for w, col in ((wi, NAVY), (wt, BLUE), (wg, TEAL), (wh, PURPLE), (wa, GREY)):
    if w is not None:
        w.sheet_properties.tabColor = col
        w.sheet_view.showGridLines = False

# =========================================================================
# TASKS – column map
# =========================================================================
NAMES = ["ID", "TASK", "OWN", "P1", "K1", "L1", "P2", "K2", "L2", "DUR", "SNET", "STAT", "ACTS", "ACTE", "PCT", "REV",
         "PS", "PE", "FS", "FE", "FD", "VAR", "DST", "SLIP", "IMP", "DEP", "CLS", "HASP", "SL1", "SL2", "EFS", "EFP",
         "XEND", "WDN"]
LT = {n: CL(i + 1) for i, n in enumerate(NAMES)}
HEADS = ["ID", "Task Name", "Owner", "Pred. 1 (ID)", "Link 1 (FS/SS/FF)", "Lag 1 (days)", "Pred. 2 (ID)",
         "Link 2 (FS/SS/FF)", "Lag 2 (days)", "Duration (work days)", "Start No Earlier Than (optional)", "Status",
         "Actual Start", "Actual End", "% Complete", "Revised Duration (if overshoot)", "Planned Start", "Planned End",
         "Forecast / Actual Start", "Forecast / Actual End", "Forecast / Actual Duration", "Duration Variance (days)",
         "Duration Status", "End Slip vs Plan (days)", "Impact (affected by)", "No. of Dependent Links",
         "class", "has pred", "pred1 slip", "pred2 slip", "eff status", "eff %", "end if no overshoot", "weighted done"]
T = lambda n: f"Tasks!${LT[n]}${FIRST}:${LT[n]}${LAST}"

# =========================================================================
# HOLIDAYS
# =========================================================================
wh.column_dimensions["A"].width = 18
wh.column_dimensions["B"].width = 38
wh.column_dimensions["C"].width = 16
wh.merge_cells("A1:C1")
put(wh, "A1", "HOLIDAY LIST", font(16, True, "FFFFFF"), PURPLE, CEN)
wh.row_dimensions[1].height = 30
wh.merge_cells("A2:C3")
put(wh, "A2", "Enter one non-working date per row (yellow cells). Weekly off day(s) are set on the Project Info sheet. "
    "Holidays move task dates and are shaded pink on the Gantt chart.", font(9, i=True, c=GREY), None, LEFT)
for col, t in zip("ABC", ("Holiday Date", "Holiday Name", "Day")):
    put(wh, f"{col}5", t, font(10, True, "FFFFFF"), NAVY, CEN, bd=BOX)
for r in range(6, 206):
    put(wh, f"A{r}", None, font(c="0000FF"), INPUT, CEN, "dd-mmm-yyyy", BOX)
    put(wh, f"B{r}", None, font(c="0000FF"), INPUT, LEFT, bd=BOX)
    put(wh, f"C{r}", f'=IF(A{r}="","",TEXT(A{r},"dddd"))', font(c=GREY), None, CEN, bd=BOX)
for i, (d, n) in enumerate(HOLIDAYS):
    wh[f"A{6+i}"].value, wh[f"B{6+i}"].value = d, n
addv(wh, DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                        errorTitle="Invalid date", error="Please enter a valid date."), "A6:A205")
wh.freeze_panes = "A6"

# =========================================================================
# PROJECT INFO
# =========================================================================
for col, w in zip("ABCDEFG", (2, 32, 30, 3, 38, 22, 14)):
    wi.column_dimensions[col].width = w
wi.column_dimensions["H"].hidden = True
wi.merge_cells("B1:G1")
put(wi, "B1", "PROJECT EXECUTION – GANTT CHART" if EXEC else "PROJECT GANTT CHART TEMPLATE", font(20, True, "FFFFFF"), NAVY, CEN)
wi.row_dimensions[1].height = 38
wi.merge_cells("B2:G2")
put(wi, "B2", "Project Info  ›  Tasks  ›  Gantt  ›  Holidays  –  everything is linked; fill the yellow cells only",
    font(10, i=True, c=GREY), None, CEN)


def section(ws, rng, text, color):
    a = rng.split(":")[0]
    ws.merge_cells(rng)
    put(ws, a, text, font(11, True, "FFFFFF"), color, LEFT)


section(wi, "B4:C4", "  PROJECT DETAILS", BLUE)
section(wi, "E4:G4", "  SCHEDULE HEALTH (auto-calculated)", TEAL)
details = [(5, "Project Name", PROJECT[0], None), (6, "Client Name", PROJECT[1], None),
           (7, "Project Manager", PROJECT[2], None), (8, "Project Code / Reference", PROJECT[3], None),
           (9, "Project Start Date", START, "dd-mmm-yyyy"), (10, "Target Completion Date", TARGET, "dd-mmm-yyyy"),
           (11, "Status (Report) Date", STATUS_DATE, "dd-mmm-yyyy")]
for r, lab, val, nf in details:
    put(wi, f"B{r}", lab, font(10, True), "EAF0FA", LEFT, bd=BOX)
    put(wi, f"C{r}", val, font(10, c="0000FF"), INPUT, CEN, nf, BOX)
wi["C9"].comment = Comment("Project execution start (EXEC: date of techno-commercially clear PO + advance).", "Template")
wi["C11"].comment = Comment("Date up to which progress is reported. =TODAY() tracks live. Un-started tasks cannot start "
                            "before this date; in-progress tasks are measured against it.", "Template")
put(wi, "B12", "Weekly Off", font(10, True), "EAF0FA", LEFT, bd=BOX)
put(wi, "C12", WEEKLY_OFF, font(10, c="0000FF"), INPUT, CEN, bd=BOX)
put(wi, "H12", '=IF(C12="Sunday only",11,IF(C12="Friday & Saturday",7,1))', font(8))
put(wi, "B13", "Gantt View (set on Gantt sheet)", font(10, True), "EAF0FA", LEFT, bd=BOX)
put(wi, "C13", "=Gantt!B2", font(10), None, CEN, bd=BOX)
addv(wi, DataValidation(type="list", formula1='"Sunday only,Saturday & Sunday,Friday & Saturday"', allow_blank=False,
                        promptTitle="Weekly off", prompt="Non-working weekday(s). Applies to every date calculation.",
                        showInputMessage=True), "C12")
addv(wi, DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                        errorTitle="Invalid date", error="Please enter a valid date (e.g. 15-Oct-2026)."), "C9:C10")

NF_DAYS = '+0" d";-0" d";0" d"'
health = [
    (5, "Planned Finish (baseline)", f'=IF(COUNT({T("PE")})=0,"",MAX({T("PE")}))', "dd-mmm-yyyy"),
    (6, "Forecast Finish (incl. delays)", f'=IF(COUNT({T("FE")})=0,"",MAX({T("FE")}))', "dd-mmm-yyyy"),
    (7, "Target Completion Date", '=IF(C10="","",C10)', "dd-mmm-yyyy"),
    (8, "Delay vs TARGET (working days)", f'=IF(OR(F6="",F7=""),"",IF(F6>=F7,{ND("F7","F6")}-1,-({ND("F6","F7")}-1)))', NF_DAYS),
    (9, "Delay vs PLAN (working days)", f'=IF(OR(F5="",F6=""),"",IF(F6>=F5,{ND("F5","F6")}-1,-({ND("F6","F5")}-1)))', NF_DAYS),
    (10, "Overall % Complete (duration weighted)", f'=IF(SUM({T("DUR")})=0,"",SUM({T("WDN")})/SUM({T("DUR")}))', "0%"),
    (11, "Calendar days late vs Target", '=IF(OR(F6="",F7=""),"",F6-F7)', NF_DAYS),
]
for r, lab, fm, nf in health:
    wi.merge_cells(f"F{r}:G{r}")
    put(wi, f"E{r}", lab, font(10, True), "E3F4F2", LEFT, bd=BOX)
    put(wi, f"F{r}", fm, font(10, True), None, CEN, nf, BOX)
for rng in ("F8", "F9", "F11"):
    wi.conditional_formatting.add(rng, Rule(type="cellIs", operator="greaterThan", formula=["0"], dxf=dxf("FFC7CE", "9C0006", True)))
    wi.conditional_formatting.add(rng, Rule(type="cellIs", operator="lessThanOrEqual", formula=["0"], dxf=dxf("C6EFCE", "006100", True)))

wi.merge_cells("B15:G15")
banner = ('=IF(F8="","Add tasks on the Tasks sheet to see project health",'
          'IF(F8>0,"⚠ PROJECT DELAYED: forecast finish "&TEXT(F6,"dd-mmm-yyyy")&" is "&F8&" working day(s) after target "&TEXT(F7,"dd-mmm-yyyy"),'
          'IF(F9>0,"⚡ AT RISK: within target, but forecast finish is "&F9&" working day(s) later than plan",'
          '"✔ ON TRACK: forecast finish "&TEXT(F6,"dd-mmm-yyyy")&" meets target "&TEXT(F7,"dd-mmm-yyyy"))))')
put(wi, "B15", banner, font(13, True, "FFFFFF"), "7F7F7F", CEN)
wi.row_dimensions[15].height = 34


def banner_cf(ws, rng, ref):
    frule(ws, rng, f'ISNUMBER(SEARCH("DELAYED",{ref}))', bg="C00000", fc="FFFFFF", bold=True, stop=True)
    frule(ws, rng, f'ISNUMBER(SEARCH("AT RISK",{ref}))', bg="ED7D31", fc="FFFFFF", bold=True, stop=True)
    frule(ws, rng, f'ISNUMBER(SEARCH("ON TRACK",{ref}))', bg="2E9E5B", fc="FFFFFF", bold=True, stop=True)


banner_cf(wi, "B15:G15", "$B$15")

# component status table
section(wi, "B17:G17", "  STATUS BY COMPONENT  (tasks whose name starts with the text in column B)", NAVY)
for col, t in zip("BCDEFG", ("Component (name prefix)", "Planned Finish", "", "Forecast Finish", "Delay vs Plan (wd)", "% Complete")):
    put(wi, f"{col}18", t, font(9, True, "FFFFFF"), BLUE, CEN, bd=BOX)
for i in range(3):
    r = 19 + i
    key = COMPONENTS[i]
    put(wi, f"B{r}", key, font(10, True, "0000FF"), INPUT, CEN, bd=BOX)
    cnt = f'COUNTIF({T("TASK")},$B{r}&"*")'
    put(wi, f"C{r}", f'=IF(OR($B{r}="",{cnt}=0),"",_xlfn.MAXIFS({T("PE")},{T("TASK")},$B{r}&"*"))', font(10, True), None, CEN, "dd-mmm-yyyy", BOX)
    put(wi, f"D{r}", None, font(), None, None, None, BOX)
    put(wi, f"E{r}", f'=IF(OR($B{r}="",{cnt}=0),"",_xlfn.MAXIFS({T("FE")},{T("TASK")},$B{r}&"*"))', font(10, True), None, CEN, "dd-mmm-yyyy", BOX)
    put(wi, f"F{r}", f'=IF(OR(C{r}="",E{r}=""),"",IF(E{r}>=C{r},{ND(f"C{r}", f"E{r}")}-1,-({ND(f"E{r}", f"C{r}")}-1)))',
        font(10, True), None, CEN, NF_DAYS, BOX)
    put(wi, f"G{r}", f'=IF(OR($B{r}="",SUMIFS({T("DUR")},{T("TASK")},$B{r}&"*")=0),"",SUMIFS({T("WDN")},{T("TASK")},$B{r}&"*")/SUMIFS({T("DUR")},{T("TASK")},$B{r}&"*"))',
        font(10, True), None, CEN, "0%", BOX)
wi.conditional_formatting.add("F19:F21", Rule(type="expression", formula=['AND(ISNUMBER(F19),F19>0)'], dxf=dxf("FFC7CE", "9C0006", True)))
wi.conditional_formatting.add("F19:F21", Rule(type="expression", formula=['AND(ISNUMBER(F19),F19<=0)'], dxf=dxf("C6EFCE", "006100", True)))

section(wi, "B23:C23", "  TASK STATISTICS", PURPLE)
section(wi, "E23:G23", "  HOW TO USE", "7F6000")
stats = [
    ("Total Tasks", f'=COUNTIF({T("TASK")},"?*")'),
    ("Completed", f'=COUNTIF({T("EFS")},"Completed")'),
    ("In Progress", f'=COUNTIF({T("EFS")},"In Progress")'),
    ("Not Started / On Hold", f'=COUNTIF({T("EFS")},"Not Started")+COUNTIF({T("EFS")},"On Hold")'),
    ("Tasks Overshooting Duration", f'=COUNTIF({T("VAR")},">0")'),
    ("Tasks Affected by Others' Delay", f'=COUNTIF({T("IMP")},"Delayed*")'),
    ("Largest Overshoot (task)", f'=IF(MAX({T("VAR")})<=0,"None",INDEX({T("TASK")},MATCH(MAX({T("VAR")}),{T("VAR")},0))&" (+"&MAX({T("VAR")})&" d)")'),
    ("Total Planned Task-Days", f'=SUM({T("DUR")})'),
]
for i, (lab, fm) in enumerate(stats):
    r = 24 + i
    put(wi, f"B{r}", lab, font(10, True), "F0E8F7", LEFT, bd=BOX)
    put(wi, f"C{r}", fm, font(10, True), None, CEN, bd=BOX)
howto = [
    "1. Fill the yellow Project Details: start / target / report date and the Weekly Off.",
    "2. Holidays sheet: list non-working dates.",
    "3. Tasks sheet: Task Name + Duration (working days) + predecessor ID(s). Link FS = after predecessor ends (default); SS = starts n days after it starts; FF = ends n days after it ends. Lag in working days.",
    "4. Update progress: Status, Actual Start/End, % Complete. Use Revised Duration when you expect an overshoot.",
    "5. Gantt sheet: pick Day / Week / Month in cell B2; scroll the timeline with C2.",
    "6. 300 task rows are pre-built; copy the last row down to add more.",
    "7. Avoid circular links (A waits for B and B waits for A).",
    "Yellow cells = input · white cells = formulas (do not overwrite).",
]
for i, t in enumerate(howto):
    r = 24 + i
    wi.merge_cells(f"E{r}:G{r}")
    put(wi, f"E{r}", t, font(9), "FFF8E1", LEFT, bd=BOX)
    wi.row_dimensions[r].height = 36 if i == 2 else 27

section(wi, "B33:G33", "  COLOUR LEGEND", NAVY)
legend = [
    (C_DONE, "FFFFFF", "Completed", "Task finished (bar portion within planned duration)"),
    (C_PROG, "FFFFFF", "In Progress", "Task started, not finished"),
    (C_NOT, "1F3864", "Not Started", "Scheduled, not yet begun"),
    (C_HOLD, "000000", "On Hold", "Task paused"),
    (C_OVER, "FFFFFF", "Overshoot", "Time used beyond the planned DURATION (task ran / is expected to run over)"),
    (C_AFF, "FFFFFF", "Delayed / Affected", "Start pushed out by a late predecessor or by the status date (see 'Impact' column)"),
    (C_GHOST, "595959", "Planned (not used)", "Planned time no longer used (finished early / started late)"),
    (C_WEEKEND, "595959", "Weekly off", "Weekly-off day (Day view)"),
    (C_HOLIDAY, "9C0057", "Holiday", "Dates from the Holidays sheet (pink header in Week/Month view = holiday inside period)"),
    (C_STATUS, "7F6000", "Status Date", "Column of the report date set above"),
    ("FFFFFF", "000000", "▶  Start", "First day of a task with no predecessor"),
    ("FFFFFF", "000000", "↳  Dependent start", "First day of a task that waits for a predecessor"),
    ("FFFFFF", "000000", "◀  End", "Last day of a task"),
    ("FFFFFF", "000000", "➜  End feeds others", "Last day of a task that other tasks depend on"),
    ("FFFFFF", "000000", "◆  Start = End", "Task starts and ends in the same period"),
]
for i, (bg, fc, lab, desc) in enumerate(legend):
    r = 34 + i
    put(wi, f"B{r}", lab, font(10, True, fc), bg, CEN, bd=BOX)
    wi.merge_cells(f"C{r}:G{r}")
    put(wi, f"C{r}", desc, font(9), None, LEFT, bd=BOX)
wi.freeze_panes = "A4"

# =========================================================================
# TASKS sheet
# =========================================================================
widths = dict(ID=5, TASK=50, OWN=14, P1=8, K1=9, L1=7, P2=8, K2=9, L2=7, DUR=10, SNET=13, STAT=13, ACTS=12, ACTE=12,
              PCT=10, REV=11, PS=12, PE=12, FS=12, FE=12, FD=10, VAR=10, DST=27, SLIP=9, IMP=28, DEP=9)
for n in NAMES:
    wt.column_dimensions[LT[n]].width = widths.get(n, 9)
wt.column_dimensions.group(LT["CLS"], LT["WDN"], hidden=True)
lastvis = LT["DEP"]

wt.merge_cells(f"A1:{lastvis}1")
put(wt, "A1", '="TASK LIST  –  "&\'Project Info\'!C5&"   |   Client: "&\'Project Info\'!C6', font(16, True, "FFFFFF"), BLUE, LEFT)
wt.row_dimensions[1].height = 32
wt.merge_cells(f"A2:{lastvis}2")
put(wt, "A2", "Type in the YELLOW cells only. Link: FS = starts after predecessor ends (blank = FS) · SS = starts n days after it "
    "starts · FF = ends n days after it ends. Lag / duration are WORKING days (weekly-off & holidays skipped).",
    font(9, i=True, c=GREY), None, LEFT)
wt.merge_cells(f"A4:{lastvis}4")
put(wt, "A4", f"={BANNER}", font(12, True, "FFFFFF"), "7F7F7F", CEN)
wt.row_dimensions[4].height = 28
banner_cf(wt, f"A4:{lastvis}4", "$A$4")
wt.merge_cells(f"A5:{lastvis}5")
put(wt, "A5", f'="Planned finish: "&IF({INFO}F5="","-",TEXT({INFO}F5,"dd-mmm-yy"))&"     |     Forecast finish: "&IF({INFO}F6="","-",TEXT({INFO}F6,"dd-mmm-yy"))'
    f'&"     |     Target: "&IF({INFO}F7="","-",TEXT({INFO}F7,"dd-mmm-yy"))&"     |     Overall progress: "&IF({INFO}F10="","-",TEXT({INFO}F10,"0%"))'
    f'&"     |     Status date: "&IF({SD}="","-",TEXT({SD},"dd-mmm-yy"))', font(10, True, NAVY), "EAF0FA", CEN)
wt.row_dimensions[5].height = 22

groups = [("ID", "SNET", "①  TASK ENTRY", BLUE), ("STAT", "REV", "②  STATUS UPDATE", "2E7D32"),
          ("PS", "DEP", "③  SCHEDULE & ANALYSIS (automatic)", PURPLE), ("CLS", "WDN", "helper", GREY)]
for a, b, t, col in groups:
    wt.merge_cells(f"{LT[a]}9:{LT[b]}9")
    put(wt, f"{LT[a]}9", t, font(11, True, "FFFFFF"), col, CEN)
for i, (n, h) in enumerate(zip(NAMES, HEADS)):
    ix = NAMES.index(n)
    col = BLUE if ix <= NAMES.index("SNET") else ("2E7D32" if ix <= NAMES.index("REV") else (PURPLE if ix <= NAMES.index("DEP") else GREY))
    put(wt, f"{LT[n]}10", h, font(9, True, "FFFFFF"), col, CEN, bd=BOX)
wt.row_dimensions[10].height = 44

HELPER_GREY = ("ID", "CLS", "HASP", "SL1", "SL2", "EFS", "EFP", "XEND", "WDN")
INPUTS = ["TASK", "OWN", "P1", "K1", "L1", "P2", "K2", "L2", "DUR", "SNET", "STAT", "ACTS", "ACTE", "PCT", "REV"]


def row_formulas(r):
    c = lambda n: f"${LT[n]}{r}"
    R = lambda n: f"${LT[n]}${FIRST}:${LT[n]}${LAST}"
    m = lambda p: f"MATCH({c(p)},{R('ID')},0)"

    def cons(p, k, l, s, e, dur):          # earliest start implied by one predecessor link
        st, en = f"INDEX({R(s)},{m(p)})", f"INDEX({R(e)},{m(p)})"
        return (f'IF({c(p)}="",0,N(IFERROR(IF({c(k)}="SS",{W(st, "N(" + c(l) + ")")},'
                f'IF({c(k)}="FF",{W(en, "N(" + c(l) + ")-" + dur + "+1")},{W(en, "1+N(" + c(l) + ")")})),0)))')

    durP = f"MAX(1,N({c('DUR')}))"
    durF = f"MAX(1,IF(ISNUMBER({c('REV')}),{c('REV')},N({c('DUR')})))"
    pl = lambda p, k, l: cons(p, k, l, "PS", "PE", durP)
    fc = lambda p, k, l: cons(p, k, l, "FS", "FE", durF)
    return {
        "ID": f'=IF({c("TASK")}="","",ROW()-{FIRST-1})',
        "PS": f'=IF(OR({c("TASK")}="",{PSD}=""),"",{W("MAX(" + PSD + ",N(" + c("SNET") + ")," + pl("P1", "K1", "L1") + "," + pl("P2", "K2", "L2") + ")-1", 1)})',
        "PE": f'=IF(OR({c("PS")}="",NOT(ISNUMBER({c("DUR")}))),"",{W(c("PS"), "MAX(1," + c("DUR") + ")-1")})',
        "FS": (f'=IF({c("PS")}="","",IF(ISNUMBER({c("ACTS")}),{c("ACTS")},'
               f'{W("MAX(" + c("PS") + "," + fc("P1", "K1", "L1") + "," + fc("P2", "K2", "L2") + ",IF(" + c("EFS") + "=" + chr(34) + "Completed" + chr(34) + ",0,N(" + SD + ")))-1", 1)}))'),
        "FD": (f'=IF(OR({c("FS")}="",NOT(ISNUMBER({c("DUR")}))),"",IF(ISNUMBER({c("ACTE")}),MAX(1,{ND(c("FS"), c("ACTE"))}),'
               f'MAX(1,IF(ISNUMBER({c("REV")}),{c("REV")},{c("DUR")}),IF(ISNUMBER({c("ACTS")}),{ND(c("ACTS"), "N(" + SD + ")")},0))))'),
        "FE": f'=IF(OR({c("FS")}="",{c("FD")}=""),"",IF(ISNUMBER({c("ACTE")}),{c("ACTE")},{W(c("FS"), c("FD") + "-1")}))',
        "VAR": f'=IF({c("FD")}="","",{c("FD")}-MAX(1,{c("DUR")}))',
        "DST": (f'=IF({c("TASK")}="","",IF({c("FD")}="","Enter Duration",IF({c("EFS")}="Completed",IF(NOT(ISNUMBER({c("ACTE")})),"Completed - enter Actual End",'
                f'IF({c("VAR")}>0,"Completed - Overshoot",IF({c("VAR")}<0,"Completed - Early","Completed - On Time"))),'
                f'IF({c("EFS")}="On Hold","On Hold",IF({c("EFS")}="In Progress",IF({c("VAR")}>0,"In Progress - Overshooting","In Progress - On Track"),'
                f'IF({c("FS")}>{c("PS")},"Delayed Start","Not Started"))))))'),
        "SLIP": (f'=IF(OR(NOT(ISNUMBER({c("PE")})),NOT(ISNUMBER({c("FE")}))),"",IF({c("FE")}>={c("PE")},{ND(c("PE"), c("FE"))}-1,'
                 f'-({ND(c("FE"), c("PE"))}-1)))'),
        "IMP": (f'=IF({c("TASK")}="","",IF(AND(MAX({c("SL1")},{c("SL2")})>0,N({c("FS")})>N({c("PS")})),'
                f'"Delayed by Task #"&IF({c("SL1")}>={c("SL2")},{c("P1")},{c("P2")})&" (+"&MAX({c("SL1")},{c("SL2")})&"d)",""))'),
        "DEP": f'=IF({c("TASK")}="","",COUNTIF({R("P1")},{c("ID")})+COUNTIF({R("P2")},{c("ID")}))',
        "CLS": (f'=IF({c("TASK")}="","",IF({c("EFS")}="Completed",1,IF({c("EFS")}="On Hold",4,IF({c("EFS")}="In Progress",2,'
                f'IF(N({c("FS")})>N({c("PS")}),7,3)))))'),
        "HASP": (f'=IF({c("TASK")}="","",IF(OR(AND({c("P1")}<>"",ISNUMBER({m("P1")})),AND({c("P2")}<>"",ISNUMBER({m("P2")}))),1,0))'),
        "SL1": f'=IF({c("P1")}="",0,N(IFERROR(INDEX({R("SLIP")},{m("P1")}),0)))',
        "SL2": f'=IF({c("P2")}="",0,N(IFERROR(INDEX({R("SLIP")},{m("P2")}),0)))',
        "EFS": (f'=IF({c("TASK")}="","",IF(ISNUMBER({c("ACTE")}),"Completed",IF({c("STAT")}="Completed","Completed",'
                f'IF({c("STAT")}<>"",{c("STAT")},IF(ISNUMBER({c("ACTS")}),"In Progress","Not Started")))))'),
        "EFP": f'=IF({c("TASK")}="","",IF({c("EFS")}="Completed",1,MIN(1,N({c("PCT")}))))',
        "XEND": f'=IF(OR({c("FS")}="",NOT(ISNUMBER({c("DUR")}))),"",{W(c("FS"), "MAX(1," + c("DUR") + ")-1")})',
        "WDN": f'=IF({c("TASK")}="","",N({c("DUR")})*{c("EFP")})',
    }


DATE_COLS = ("SNET", "ACTS", "ACTE", "PS", "PE", "FS", "FE", "XEND")
for r in range(FIRST, LAST + 1):
    for n, fm in row_formulas(r).items():
        cell = wt[f"{LT[n]}{r}"]
        cell.value = fm
        cell.font = font(9, c=GREY if n in HELPER_GREY else "000000")
        cell.border = BOX
        cell.alignment = LEFT if n in ("DST", "IMP") else CEN
    for n in INPUTS:
        cell = wt[f"{LT[n]}{r}"]
        cell.fill, cell.border, cell.font = fill(INPUT), BOX, font(10, c="0000FF")
        cell.alignment = LEFT if n in ("TASK", "OWN") else CEN
    for n in DATE_COLS:
        wt[f"{LT[n]}{r}"].number_format = "dd-mmm-yy"
    wt[f"{LT['PCT']}{r}"].number_format = "0%"
    wt[f"{LT['EFP']}{r}"].number_format = "0%"
    for n in ("VAR", "SLIP"):
        wt[f"{LT[n]}{r}"].number_format = "+0;-0;0"
    wt.row_dimensions[r].height = 18

for i, t in enumerate(EX_TASKS):
    r = FIRST + i
    vals = {"TASK": t["name"], "OWN": t["owner"], "DUR": t["dur"], "SNET": t.get("snet"),
            "STAT": t.get("status", "Not Started"), "ACTS": t.get("acts"), "ACTE": t.get("acte"),
            "PCT": t.get("pct"), "REV": t.get("rev")}
    for slot, p in (("1", t["p1"]), ("2", t["p2"])):
        if p:
            vals["P" + slot], vals["K" + slot], vals["L" + slot] = p[0], p[1], p[2]
    for n, v in vals.items():
        if v is not None:
            wt[f"{LT[n]}{r}"].value = v

rngc = lambda n: f"{LT[n]}{FIRST}:{LT[n]}{LAST}"
addv(wt, DataValidation(type="list", formula1='"Not Started,In Progress,Completed,On Hold"', allow_blank=True,
                        promptTitle="Status", prompt="Pick the current task status.", showInputMessage=True), rngc("STAT"))
for p, k, l in (("P1", "K1", "L1"), ("P2", "K2", "L2")):
    addv(wt, DataValidation(type="whole", operator="between", formula1="1", formula2="9999", allow_blank=True,
                            promptTitle="Predecessor", prompt="ID (first column) of the task this one depends on. Leave blank if none.",
                            showInputMessage=True, errorTitle="Invalid", error="Enter a task ID number."), rngc(p))
    addv(wt, DataValidation(type="list", formula1='"FS,SS,FF"', allow_blank=True, promptTitle="Link type",
                            prompt="FS = start after predecessor ends (default). SS = start n days after it starts. FF = end n days after it ends.",
                            showInputMessage=True), rngc(k))
    addv(wt, DataValidation(type="whole", operator="between", formula1="-365", formula2="365", allow_blank=True,
                            promptTitle="Lag", prompt="Working days added to the link (negative = overlap).",
                            showInputMessage=True), rngc(l))
addv(wt, DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="1", allow_blank=True,
                        promptTitle="Duration", prompt="Planned duration in WORKING days (min 1).", showInputMessage=True,
                        errorTitle="Invalid", error="Enter a number of days >= 1."), rngc("DUR"))
addv(wt, DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="1", allow_blank=True,
                        promptTitle="Revised duration", prompt="New total working days if you now expect the task to take longer/shorter.",
                        showInputMessage=True), rngc("REV"))
addv(wt, DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                        errorTitle="Invalid date", error="Enter a valid date."), f"{LT['SNET']}{FIRST}:{LT['SNET']}{LAST}")
addv(wt, DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                        errorTitle="Invalid date", error="Enter a valid date."), f"{LT['ACTS']}{FIRST}:{LT['ACTE']}{LAST}")
addv(wt, DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                        promptTitle="% Complete", prompt="Enter 0% to 100%.", showInputMessage=True,
                        errorTitle="Invalid", error="Enter a value between 0% and 100%."), rngc("PCT"))

V, I_, D_, S_ = LT["VAR"], LT["IMP"], LT["DST"], LT["SLIP"]
frule(wt, f"A{FIRST}:B{LAST}", f"AND(ISNUMBER(${V}{FIRST}),${V}{FIRST}>0)", bg="FFC7CE", fc="9C0006", bold=True)
frule(wt, f"A{FIRST}:B{LAST}", f'${I_}{FIRST}<>""', bg="FCE4D6", fc="C55A11", bold=True)
for txt, bg, fc in (("Completed - On Time", "C6EFCE", "006100"), ("Completed - Early", "B7E1E8", "0B5563"),
                    ("Completed - Overshoot", "FFC7CE", "9C0006"), ("In Progress - On Track", "DDEBF7", "1F4E79"),
                    ("In Progress - Overshooting", "FF9999", "7F0000"), ("Delayed Start", "FFD8B0", "974706"),
                    ("On Hold", "FFF2CC", "7F6000"), ("Not Started", "EDEDED", "595959"), ("Enter", "FFEB9C", "9C5700")):
    frule(wt, f"{D_}{FIRST}:{D_}{LAST}", f'LEFT(${D_}{FIRST},{len(txt)})="{txt}"', bg=bg, fc=fc, bold=True, stop=True)
for col in (V, S_):
    frule(wt, f"{col}{FIRST}:{col}{LAST}", f"AND(ISNUMBER(${col}{FIRST}),${col}{FIRST}>0)", bg="FFC7CE", fc="9C0006", bold=True)
    frule(wt, f"{col}{FIRST}:{col}{LAST}", f"AND(ISNUMBER(${col}{FIRST}),${col}{FIRST}<0)", bg="C6EFCE", fc="006100", bold=True)
frule(wt, f"{I_}{FIRST}:{I_}{LAST}", f'${I_}{FIRST}<>""', bg="FFD8B0", fc="974706", bold=True)
st = LT["STAT"]
for txt, bg, fc in (("Completed", "C6EFCE", "006100"), ("In Progress", "DDEBF7", "1F4E79"), ("On Hold", "FFF2CC", "7F6000")):
    frule(wt, f"{st}{FIRST}:{st}{LAST}", f'${st}{FIRST}="{txt}"', bg=bg, fc=fc, bold=True)
wt.freeze_panes = f"C{FIRST}"
wt.auto_filter.ref = f"A10:{lastvis}{LAST}"

# =========================================================================
# GANTT
# =========================================================================
for col, w in dict(A=5, B=50, C=15, D=10, E=10, F=10, G=10, H=12, I=7, J=8, K=3, L=3, M=3, N=3).items():
    wg.column_dimensions[col].width = w
wg.column_dimensions.group("K", "N", hidden=True)
for c in range(TL0, TLN + 1):
    wg.column_dimensions[CL(c)].width = 4.4
for r in (5, 6, 7):
    wg.row_dimensions[r].hidden = True
wg.row_dimensions[1].height = 30
wg.row_dimensions[8].height = 26
wg.row_dimensions[10].height = 30

wg.merge_cells("A1:J1")
put(wg, "A1", '="GANTT  |  "&\'Project Info\'!C5&"  |  "&\'Project Info\'!C6', font(14, True, "FFFFFF"), NAVY, LEFT)
bn = TL0 + 50
wg.merge_cells(start_row=1, start_column=TL0, end_row=1, end_column=bn)
put(wg, f"{CL(TL0)}1", f"={BANNER}", font(12, True, "FFFFFF"), "7F7F7F", CEN)
banner_cf(wg, f"{CL(TL0)}1:{CL(bn)}1", f"${CL(TL0)}$1")

put(wg, "B2", "Week", font(11, True, "0000FF"), "FFE699", CEN, '"VIEW ▸  "@', BOX)
wg.merge_cells("C2:D2")
put(wg, "C2", 0, font(10, True, "0000FF"), "FFE699", CEN, '"Scroll "+0;"Scroll "-0;"Scroll 0"', BOX)
wg.merge_cells("E2:J2")
put(wg, "E2", "◀ pick Day / Week / Month   ·   Scroll = shift timeline by n periods (+ later / − earlier)", font(8, i=True, c=GREY), None, LEFT)
wg.row_dimensions[2].height = 24
addv(wg, DataValidation(type="list", formula1='"Day,Week,Month"', allow_blank=False, showDropDown=False,
                        promptTitle="Gantt view", prompt="Choose Day, Week or Month.", showInputMessage=True,
                        errorTitle="Invalid", error="Choose Day, Week or Month."), "B2")
addv(wg, DataValidation(type="whole", operator="between", formula1="-500", formula2="500", allow_blank=False,
                        promptTitle="Scroll", prompt="Whole number of periods to shift the timeline (negative = earlier).",
                        showInputMessage=True), "C2")

wg.merge_cells("A3:J3")
put(wg, "A3", "LEGEND  ▸  bar colour = task status · RED part of a bar = time over the planned duration", font(8, True, GREY), None, LEFT)
wg.merge_cells("A4:J4")
put(wg, "A4", "Orange task name = delayed by predecessor · Red task name = duration overshoot · Pink header = holiday", font(8, True, GREY), None, LEFT)
leg_a = [("Completed", C_DONE, "FFFFFF"), ("In Progress", C_PROG, "FFFFFF"), ("Not Started", C_NOT, NAVY),
         ("On Hold", C_HOLD, "000000"), ("Overshoot (over duration)", C_OVER, "FFFFFF"),
         ("Delayed / Affected", C_AFF, "FFFFFF"), ("Planned, unused", C_GHOST, GREY)]
leg_b = [("Weekly off", C_WEEKEND, GREY), ("Holiday", C_HOLIDAY, "9C0057"), ("Status date", C_STATUS, "7F6000"),
         ("▶ Start", "F2F2F2", "000000"), ("↳ Starts after pred.", "F2F2F2", "000000"), ("◀ End", "F2F2F2", "000000"),
         ("➜ End feeds others", "F2F2F2", "000000"), ("◆ Start = End", "F2F2F2", "000000")]
for row, items in ((3, leg_a), (4, leg_b)):
    for i, (t, bg, fc) in enumerate(items):
        c0 = TL0 + i * 8
        wg.merge_cells(start_row=row, start_column=c0, end_row=row, end_column=c0 + 7)
        put(wg, f"{CL(c0)}{row}", t, font(8, True, fc), bg, CEN)
        for cc in range(c0, c0 + 8):
            wg.cell(row, cc).border = BOX

wg.merge_cells("A8:J9")
put(wg, "A8", "TASK SCHEDULE   (Plan vs Forecast / Actual)", font(11, True, "FFFFFF"), BLUE, CEN)
for i, h in enumerate(["ID", "Task", "Depends On", "Plan Start", "Plan End", "Fcst / Actual Start",
                       "Fcst / Actual End", "Status", "% Done", "Var (days)"], start=1):
    put(wg, f"{CL(i)}10", h, font(9, True, "FFFFFF"), NAVY, CEN, bd=BOX)
put(wg, "K5", f'=IF($B$2="Day",{PSD},IF($B$2="Week",{PSD}-WEEKDAY({PSD},2)+1,DATE(YEAR({PSD}),MONTH({PSD}),1)))', font(8))
wg["K5"].number_format = "dd-mmm-yy"

for k in range(NCOL):
    c = TL0 + k
    X, Wp = CL(c), CL(c - 1)
    wg[f"{X}5"] = f'=IF($B$2="Day",$K$5+{k}+$C$2,IF($B$2="Week",$K$5+7*({k}+$C$2),EDATE($K$5,{k}+$C$2)))'
    wg[f"{X}6"] = f'=IF($B$2="Day",{X}5,IF($B$2="Week",{X}5+6,EOMONTH({X}5,0)))'
    wg[f"{X}7"] = (f'=IF($B$2="Day",IF(COUNTIF({HOL},{X}5)>0,2,IF(NETWORKDAYS.INTL({X}5,{X}5,{WKN})=0,1,0)),'
                   f'IF(COUNTIFS({HOL},">="&{X}5,{HOL},"<="&{X}6)>0,3,0))+IF(AND({X}5<={SD},{X}6>={SD}),10,0)')
    if k == 0:
        wg[f"{X}8"] = f'=IF($B$2="Month",TEXT({X}5,"yyyy"),TEXT({X}5,"mmm")&CHAR(10)&TEXT({X}5,"yy"))'
    else:
        wg[f"{X}8"] = (f'=IF($B$2="Month",IF(YEAR({X}5)<>YEAR({Wp}5),TEXT({X}5,"yyyy"),""),'
                       f'IF(MONTH({X}5)<>MONTH({Wp}5),TEXT({X}5,"mmm")&CHAR(10)&TEXT({X}5,"yy"),""))')
    wg[f"{X}9"] = f'=IF($B$2="Day",TEXT({X}5,"dd"),IF($B$2="Week","W"&TEXT(WEEKNUM({X}5,21),"00"),TEXT({X}5,"mmm")))'
    wg[f"{X}10"] = f'=IF($B$2="Day",LEFT(TEXT({X}5,"ddd"),1),IF($B$2="Week",TEXT({X}5,"dd"),""))'
    for r in (8, 9, 10):
        cell = wg[f"{X}{r}"]
        cell.font = font(7 if r != 9 else 8, True, "FFFFFF")
        cell.fill = fill(NAVY if r != 9 else BLUE)
        cell.alignment = CEN
        cell.border = BOX
    wg[f"{X}5"].number_format = wg[f"{X}6"].number_format = "dd-mmm-yy"

hair = Side(style="hair", color="D0D0D0")
for r in range(FIRST, LAST + 1):
    tk = lambda n: f"Tasks!${LT[n]}{r}"

    def ptxt(p, k, l):
        return (f'IF({tk(p)}="","",{tk(p)}&IF(OR({tk(k)}="",{tk(k)}="FS"),IF(N({tk(l)})=0,"","FS"),{tk(k)})'
                f'&IF(N({tk(l)})=0,"",IF({tk(l)}>0,"+","")&{tk(l)}))')

    L_ = {
        "A": f'=IF({tk("TASK")}="","",{tk("ID")})',
        "B": f'=IF({tk("TASK")}="","",{tk("TASK")})',
        "C": f'=IF($B{r}="","",IF(AND({tk("P1")}="",{tk("P2")}=""),"-",{ptxt("P1", "K1", "L1")}&IF(AND({tk("P1")}<>"",{tk("P2")}<>""),", ","")&{ptxt("P2", "K2", "L2")}))',
        "D": f'=IF(OR($B{r}="",{tk("PS")}=""),"",{tk("PS")})',
        "E": f'=IF(OR($B{r}="",{tk("PE")}=""),"",{tk("PE")})',
        "F": f'=IF(OR($B{r}="",{tk("FS")}=""),"",{tk("FS")})',
        "G": f'=IF(OR($B{r}="",{tk("FE")}=""),"",{tk("FE")})',
        "H": f'=IF($B{r}="","",{tk("EFS")})',
        "I": f'=IF($B{r}="","",{tk("EFP")})',
        "J": f'=IF(OR($B{r}="",{tk("VAR")}=""),"",{tk("VAR")})',
        "K": f'=IF($B{r}="","",{tk("CLS")})',
        "L": f'=IF($B{r}="","",{tk("HASP")})',
        "M": f'=IF($B{r}="","",{tk("DEP")})',
        "N": f'=IF(OR($B{r}="",{tk("XEND")}=""),"",{tk("XEND")})',
    }
    for col, fm in L_.items():
        cell = wg[f"{col}{r}"]
        cell.value = fm
        cell.font = font(9, True if col == "B" else False)
        cell.border = BOX
        cell.alignment = LEFT if col in "BC" else CEN
        if col in "DEFG":
            cell.number_format = "dd-mmm-yy"
    wg[f"I{r}"].number_format = "0%"
    wg[f"J{r}"].number_format = "+0;-0;0"
    wg.row_dimensions[r].height = 18
    for k in range(NCOL):
        X = CL(TL0 + k)
        fm = (f'=IF(OR($A{r}="",NOT(ISNUMBER($G{r}))),"",IF(AND($F{r}<={X}$6,$G{r}>={X}$5),'
              f'IF({X}$5>$N{r},5,$K{r})*10+IF(AND($F{r}>={X}$5,$F{r}<={X}$6),'
              f'IF(AND($G{r}>={X}$5,$G{r}<={X}$6),5,IF($L{r}=1,2,1)),'
              f'IF(AND($G{r}>={X}$5,$G{r}<={X}$6),IF($M{r}>0,4,3),0)),'
              f'IF(AND($D{r}<={X}$6,$E{r}>={X}$5),60,0)))')
        cell = wg[f"{X}{r}"]
        cell.value = fm
        cell.font = font(8, True, "FFFFFF")
        cell.alignment = CEN
        cell.number_format = ";;;"
        cell.border = Border(left=hair, right=hair, top=hair, bottom=hair)

tl = f"{CL(TL0)}{FIRST}:{CL(TLN)}{LAST}"
cls_col = {1: (C_DONE, "FFFFFF"), 2: (C_PROG, "FFFFFF"), 3: (C_NOT, NAVY), 4: (C_HOLD, "000000"),
           5: (C_OVER, "FFFFFF"), 7: (C_AFF, "FFFFFF")}
marks = {1: "▶", 2: "↳", 3: "◀", 4: "➜", 5: "◆"}
nfid = 200
for k_, (bg, fc) in cls_col.items():
    for m in range(6):
        d = dxf(bg, fc, nf=";;;", nfid=nfid) if m == 0 else dxf(bg, fc if k_ != 3 else "000000", True, nf=f'"{marks[m]}"', nfid=nfid)
        nfid += 1
        wg.conditional_formatting.add(tl, Rule(type="cellIs", operator="equal", formula=[str(k_ * 10 + m)], dxf=d, stopIfTrue=True))
wg.conditional_formatting.add(tl, Rule(type="cellIs", operator="equal", formula=["60"],
                                       dxf=dxf(C_GHOST, C_GHOST, nf=";;;", nfid=nfid), stopIfTrue=True))
f7 = f"{CL(TL0)}$7"
frule(wg, tl, f"{f7}>=10", bg=C_STATUS, stop=True)
frule(wg, tl, f"MOD({f7},10)=2", bg=C_HOLIDAY, stop=True)
frule(wg, tl, f"MOD({f7},10)=1", bg=C_WEEKEND, stop=True)
hdr = f"{CL(TL0)}8:{CL(TLN)}10"
frule(wg, hdr, f"{f7}>=10", bg="F57C00", fc="FFFFFF", bold=True, stop=True)
frule(wg, hdr, f"OR(MOD({f7},10)=2,MOD({f7},10)=3)", bg="C2185B", fc="FFFFFF", bold=True, stop=True)
frule(wg, hdr, f"MOD({f7},10)=1", bg="7F7F7F", fc="FFFFFF", stop=True)

wg.conditional_formatting.add(f"A{FIRST}:B{LAST}", Rule(type="expression", formula=[f'AND(ISNUMBER($J{FIRST}),$J{FIRST}>0)'],
                                                       dxf=dxf("FFC7CE", "9C0006", True), stopIfTrue=True))
wg.conditional_formatting.add(f"A{FIRST}:B{LAST}", Rule(type="expression", formula=[f'$K{FIRST}=7'],
                                                       dxf=dxf("FCE4D6", "C55A11", True), stopIfTrue=True))
for txt, bg, fc in (("Completed", "C6EFCE", "006100"), ("In Progress", "DDEBF7", "1F4E79"),
                    ("On Hold", "FFF2CC", "7F6000"), ("Not Started", "EDEDED", "595959")):
    frule(wg, f"H{FIRST}:H{LAST}", f'$H{FIRST}="{txt}"', bg=bg, fc=fc, bold=True)
frule(wg, f"J{FIRST}:J{LAST}", f"AND(ISNUMBER($J{FIRST}),$J{FIRST}>0)", bg="FFC7CE", fc="9C0006", bold=True)
frule(wg, f"J{FIRST}:J{LAST}", f"AND(ISNUMBER($J{FIRST}),$J{FIRST}<0)", bg="C6EFCE", fc="006100", bold=True)
wg.conditional_formatting.add(f"I{FIRST}:I{LAST}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                                              end_value=1, color="63BE7B", showValue=True))
wg.freeze_panes = f"C{FIRST}"

# =========================================================================
# ASSUMPTIONS (execution workbook only)
# =========================================================================
if EXEC:
    wa.column_dimensions["A"].width = 6
    wa.column_dimensions["B"].width = 56
    wa.column_dimensions["C"].width = 10
    wa.column_dimensions["D"].width = 110
    wa.merge_cells("A1:D1")
    put(wa, "A1", "ASSUMPTIONS & BASIS OF DURATIONS", font(16, True, "FFFFFF"), GREY, LEFT)
    wa.row_dimensions[1].height = 30
    general = [
        "Project start = 15-Oct-2026 (date of techno-commercially clear PO + advance), as given.",
        "Working calendar: Monday–Saturday, Sunday off (Project Info › Weekly Off). All durations are working days.",
        "Holidays: Dussehra 20-Oct-2026; Diwali 4 days (Sat 7, Mon 9, Tue 10, Wed 11 Nov 2026 – Sun 8 Nov is already off; CONFIRM dates); New Year 1-Jan-2027.",
        "Quantities: 500 Reflectors, 500 Feed assemblies (3,000 components = 6 per set), 500 Metal Structure sets.",
        "Rates (per day): Reflector production 35, finishing & packing 35. Feed: molding 400 (2 machines × 200), machining 50 (2 machines × 25), "
        "finishing 60 (2 stations × 30, unchanged from earlier input), assembly & test 30, packing 100. Metal: assembly & packing 75.",
        "Vehicles: Reflector 70, Feed 150, Metal 50 per vehicle. Dispatch is rolling – a vehicle leaves when a full load is packed.",
        "Reflector finishing starts the working day after production ends (no overlap; ≥12 h cooling gap).",
        "Feed: molding & machining at site A; after each 250 sets the lot goes to site B (1 working day transport) for finishing, assembly, testing, packing, dispatch.",
        "Feed machining order: round-1 components (both batches) first, then round-2, so machines stay busy while round-2 molds are being made.",
        "Link types: FS (finish-to-start), SS+n (starts n days after predecessor starts), FF+n (ends n days after predecessor ends).",
        "Target Completion Date is set to the baseline planned finish (23-Apr-2027); overwrite with the contractual date.",
        "Status (Report) Date = TODAY(): un-started tasks cannot be forecast to start before today, so unreported progress shows as delay.",
        "Owner names, client and project manager are placeholders – edit on Project Info / Tasks.",
    ]
    wa["A3"].value = "General assumptions"
    wa["A3"].font = font(11, True, NAVY)
    for i, t in enumerate(general):
        r = 4 + i
        wa.merge_cells(f"A{r}:D{r}")
        put(wa, f"A{r}", f"{i+1}.  {t}", font(10), None, LEFT)
        wa.row_dimensions[r].height = 30 if len(t) > 120 else 18
    r0 = 4 + len(general) + 1
    for col, t in zip("ABCD", ("ID", "Task", "Days", "Basis of duration / link")):
        put(wa, f"{col}{r0}", t, font(10, True, "FFFFFF"), NAVY, CEN, bd=BOX)
    for i, (n, o, p1, p2, d, basis) in enumerate(TASKS):
        r = r0 + 1 + i
        put(wa, f"A{r}", i + 1, font(10), None, CEN, bd=BOX)
        put(wa, f"B{r}", n, font(10), None, LEFT, bd=BOX)
        put(wa, f"C{r}", d, font(10), None, CEN, bd=BOX)
        put(wa, f"D{r}", basis, font(9), None, LEFT, bd=BOX)
        wa.row_dimensions[r].height = 40 if len(basis) > 120 else 26 if len(basis) > 70 else 18
    wa.freeze_panes = f"A{r0+1}"

for ws in (wt, wg, wi):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
